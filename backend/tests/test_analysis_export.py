import json
import os
from io import BytesIO

os.environ["ENV"] = "dev"
os.environ["JWT_SECRET_KEY"] = "test-secret-key-with-at-least-32-characters"
os.environ["DATABASE_URL"] = "sqlite:///:memory:"

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.core.security import create_access_token
from app.main import app
from app.models.analysis import Analysis
from app.models.conversation import Conversation
from app.models.dataset import Dataset
from app.models.user import User

test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestSession = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)
client = TestClient(app)


def override_get_db():
    db = TestSession()
    try:
        yield db
    finally:
        db.close()


def setup_function():
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    app.dependency_overrides[get_db] = override_get_db


def teardown_function():
    app.dependency_overrides.clear()


def _create_analysis() -> tuple[int, int, int]:
    with TestSession() as db:
        user = User(
            name="Exporter",
            email="export@example.com",
            password_hash="unused",
            is_verified=True,
        )
        other_user = User(
            name="Other",
            email="other-export@example.com",
            password_hash="unused",
            is_verified=True,
        )
        db.add_all([user, other_user])
        db.flush()
        dataset = Dataset(
            user_id=user.id,
            filename="ventes.csv",
            file_path="unused.csv",
            file_type="csv",
            row_count=1,
            column_count=2,
        )
        db.add(dataset)
        db.flush()
        analysis = Analysis(
            user_id=user.id,
            dataset_id=dataset.id,
            question="Question <script>alert(1)</script>",
            answer="+SUM(A1:A2)",
            result_type="text",
            result_data=json.dumps({
                "dataset_ids": [dataset.id],
                "result_table": {
                    "columns": ["région", "montant"],
                    "rows": [["=2+2", 42]],
                },
            }),
            chart_config=json.dumps({
                "type": "bar",
                "title": "Montants par région",
                "x_label": "Région",
                "y_label": "Montant",
                "labels": ["Nord", "Sud"],
                "series": [{"name": "Ventes", "values": [10, 42]}],
            }),
        )
        db.add(analysis)
        db.commit()
        return user.id, other_user.id, analysis.id


def _headers(user_id: int) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(str(user_id))}"}


def test_xlsx_export_contains_sheets_safe_cells_and_native_chart():
    user_id, _, analysis_id = _create_analysis()

    response = client.get(
        f"/analysis/{analysis_id}/export?format=xlsx",
        headers=_headers(user_id),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == 'attachment; filename="analyse-1-resultat.xlsx"'
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Résumé", "Données"]
    assert workbook["Résumé"]["B3"].value == "'+SUM(A1:A2)"
    assert workbook["Résumé"]["B5"].value == "ventes.csv"
    assert workbook["Données"]["A2"].value == "'=2+2"
    assert workbook["Données"]["B2"].value == 42
    assert len(workbook["Données"]._charts) == 1


def test_pdf_export_is_readable_and_escapes_user_text():
    user_id, _, analysis_id = _create_analysis()

    response = client.get(
        f"/analysis/{analysis_id}/export?format=pdf",
        headers=_headers(user_id),
    )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-disposition"] == 'attachment; filename="analyse-question-script-alert-1-script.pdf"'
    pdf = fitz.open(stream=response.content, filetype="pdf")
    content = "\n".join(page.get_text() for page in pdf)
    assert "AI Data Analyst" in content
    assert "Rapport d’analyse" in content
    assert "Question <script>alert(1)</script>" in content
    assert "+SUM(A1:A2)" in content
    assert "ventes.csv" in content
    assert len(pdf[0].get_images()) >= 2
    assert "région" in content


def test_pdf_conversation_export_contains_all_turns_and_uses_main_question_filename():
    with TestSession() as db:
        user = User(name="User", email="conversation-pdf@example.com", password_hash="unused", is_verified=True)
        db.add(user)
        db.flush()
        dataset = Dataset(
            user_id=user.id,
            filename="ventes.csv",
            file_path="unused.csv",
            file_type="csv",
            row_count=1,
            column_count=2,
        )
        db.add(dataset)
        db.flush()
        conversation = Conversation(
            user_id=user.id,
            title="Quel est le meilleur produit ?",
            dataset_ids=json.dumps([dataset.id]),
        )
        db.add(conversation)
        db.flush()
        first = Analysis(
            user_id=user.id,
            conversation_id=conversation.id,
            dataset_id=dataset.id,
            question="Quel est le meilleur produit ?",
            answer="Le meilleur produit est Premium.",
            result_type="text",
            result_data=json.dumps({
                "dataset_ids": [dataset.id],
                "result_table": {"columns": ["produit"], "rows": [["Premium"]]},
            }),
            chart_config=json.dumps({
                "type": "bar",
                "title": "Ventes par produit",
                "labels": ["Premium"],
                "series": [{"name": "Ventes", "values": [42]}],
            }),
        )
        second = Analysis(
            user_id=user.id,
            conversation_id=conversation.id,
            dataset_id=dataset.id,
            question="Et en Sud ?",
            answer="Premium reste en tête dans le Sud.",
            result_type="text",
            result_data=json.dumps({
                "dataset_ids": [dataset.id],
                "result_table": {"columns": ["région", "produit"], "rows": [["Sud", "Premium"]]},
            }),
        )
        db.add_all([first, second])
        db.commit()
        user_id, latest_analysis_id = user.id, second.id

    response = client.get(
        f"/analysis/{latest_analysis_id}/export?format=pdf",
        headers=_headers(user_id),
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="analyse-meilleur-produit.pdf"'
    pdf = fitz.open(stream=response.content, filetype="pdf")
    text = "\n".join(page.get_text() for page in pdf)
    assert "Quel est le meilleur produit ?" in text
    assert "Et en Sud ?" in text
    assert "Premium reste en tête dans le Sud." in text
    assert "Tableau de résultat" in text
    assert sum(len(page.get_images()) for page in pdf) >= 2


def test_pdf_filename_adds_id_when_user_has_a_duplicate_question():
    user_id, _, analysis_id = _create_analysis()
    with TestSession() as db:
        analysis = db.get(Analysis, analysis_id)
        db.add(Analysis(
            user_id=user_id,
            dataset_id=analysis.dataset_id,
            question=analysis.question,
            answer="Une autre réponse.",
            result_type="text",
        ))
        db.commit()

    response = client.get(
        f"/analysis/{analysis_id}/export?format=pdf",
        headers=_headers(user_id),
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == f'attachment; filename="analyse-question-script-alert-1-script-{analysis_id}.pdf"'


def test_export_is_not_available_to_another_user():
    _, other_user_id, analysis_id = _create_analysis()

    response = client.get(
        f"/analysis/{analysis_id}/export?format=pdf",
        headers=_headers(other_user_id),
    )

    assert response.status_code == 404
