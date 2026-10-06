from __future__ import annotations

import io
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import matplotlib
from matplotlib import font_manager

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


_font_path = font_manager.findfont("DejaVu Sans")
_bold_font_path = font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight="bold"))
_logo_path = Path(__file__).resolve().parents[3] / "ai-data-analyst-icon.jpg"
pdfmetrics.registerFont(TTFont("ExportSans", _font_path))
pdfmetrics.registerFont(TTFont("ExportSans-Bold", _bold_font_path))
pdfmetrics.registerFontFamily(
    "ExportSans",
    normal="ExportSans",
    bold="ExportSans-Bold",
    italic="ExportSans",
    boldItalic="ExportSans-Bold",
)


def _clean_table(table: Any) -> tuple[list[str], list[list[Any]]]:
    if not isinstance(table, dict):
        return [], []
    columns = table.get("columns")
    rows = table.get("rows")
    if not isinstance(columns, list) or not isinstance(rows, list):
        return [], []
    safe_columns = [str(value) for value in columns]
    safe_rows = [
        [row[index] if index < len(row) else None for index in range(len(safe_columns))]
        for row in rows[:1000]
        if isinstance(row, list)
    ]
    return safe_columns, safe_rows


def _excel_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, (date, datetime, int, float, bool)) or value is None:
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    return str(value)


def _normalize_chart(config: Any) -> dict[str, Any] | None:
    if not isinstance(config, dict):
        return None
    labels = config.get("labels")
    if not isinstance(labels, list) or not labels:
        return None
    series = config.get("series")
    if not isinstance(series, list) and isinstance(config.get("values"), list):
        series = [{"name": config.get("title") or "Valeur", "values": config["values"]}]
    if not isinstance(series, list) or not series:
        return None
    chart_type = config.get("type", "bar")
    if chart_type not in {"bar", "line", "pie", "scatter"}:
        return None
    normalized_series = []
    for item in series:
        if not isinstance(item, dict):
            return None
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(labels):
            return None
        if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) for value in values):
            return None
        normalized_series.append({
            "name": str(item.get("name") or f"Série {len(normalized_series) + 1}"),
            "values": values,
        })
    if chart_type == "pie" and (len(normalized_series) != 1 or len(labels) > 8):
        return None
    if chart_type == "scatter":
        try:
            [float(label) for label in labels]
        except (TypeError, ValueError):
            return None
    return {
        **config,
        "type": chart_type,
        "labels": [str(label) for label in labels],
        "series": normalized_series,
    }


def _chart_from_config(config: Any, worksheet, anchor: str):
    config = _normalize_chart(config)
    if config is None:
        return
    chart_type = config["type"]
    labels = config["labels"]
    series = config["series"]

    first_column = worksheet.max_column + 2
    worksheet.cell(row=1, column=first_column, value="Catégorie")
    for row_index, label in enumerate(labels, start=2):
        worksheet.cell(row=row_index, column=first_column, value=_excel_safe(label))
    for series_index, item in enumerate(series, start=1):
        if not isinstance(item, dict) or not isinstance(item.get("values"), list):
            return
        if len(item["values"]) != len(labels):
            return
        column = first_column + series_index
        worksheet.cell(row=1, column=column, value=_excel_safe(item.get("name", f"Série {series_index}")))
        for row_index, value in enumerate(item["values"], start=2):
            worksheet.cell(row=row_index, column=column, value=_excel_safe(value))

    if chart_type == "line":
        chart = LineChart()
    elif chart_type == "pie" and len(series) == 1:
        chart = PieChart()
    elif chart_type == "scatter":
        try:
            chart = ScatterChart()
            [float(label) for label in labels]
        except (TypeError, ValueError):
            return
        for series_index, item in enumerate(series, start=1):
            x_reference = Reference(worksheet, min_col=first_column, min_row=2, max_row=len(labels) + 1)
            y_reference = Reference(worksheet, min_col=first_column + series_index, min_row=2, max_row=len(labels) + 1)
            chart.series.append(Series(y_reference, x_reference, title=str(item.get("name", ""))))
    else:
        chart = BarChart()

    if chart_type != "scatter":
        data = Reference(
            worksheet,
            min_col=first_column + 1,
            max_col=first_column + len(series),
            min_row=1,
            max_row=len(labels) + 1,
        )
        categories = Reference(worksheet, min_col=first_column, min_row=2, max_row=len(labels) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
    chart.title = str(config.get("title") or "Graphique")
    if chart_type != "pie":
        chart.x_axis.title = str(config.get("x_label") or "")
        chart.y_axis.title = str(config.get("y_label") or "")
    worksheet.add_chart(chart, anchor)


def build_xlsx_export(
    *,
    question: str,
    answer: str,
    created_at: datetime,
    source_files: list[str],
    result_table: Any,
    chart_config: Any,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Résumé"
    summary.append(["Élément", "Valeur"])
    summary.append(["Question", _excel_safe(question)])
    summary.append(["Réponse", _excel_safe(answer)])
    summary.append(["Date", _excel_safe(created_at)])
    summary.append(["Fichiers sources", _excel_safe(", ".join(source_files))])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 90
    summary.freeze_panes = "A2"

    data_sheet = workbook.create_sheet("Données")
    columns, rows = _clean_table(result_table)
    if columns:
        data_sheet.append([_excel_safe(value) for value in columns])
        for row in rows:
            data_sheet.append([_excel_safe(value) for value in row])
        data_sheet.freeze_panes = "A2"
        for index, column in enumerate(columns, start=1):
            data_sheet.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 12), 36)
    else:
        data_sheet.append(["Aucun tableau de résultat disponible."])

    _chart_from_config(chart_config, data_sheet, "J2")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _chart_png(config: Any) -> bytes | None:
    config = _normalize_chart(config)
    if config is None:
        return None
    labels = config["labels"]
    series = config["series"]

    figure, axis = plt.subplots(figsize=(8, 4.5), dpi=120)
    chart_type = config.get("type", "bar")
    if chart_type == "pie" and len(series) == 1:
        axis.pie(series[0]["values"], labels=labels, autopct="%1.1f%%")
    elif chart_type == "scatter":
        for item in series:
            axis.scatter([float(label) for label in labels], item["values"], label=item["name"])
        axis.set_xlabel(config.get("x_label") or "")
        axis.set_ylabel(config.get("y_label") or "")
        axis.legend()
    else:
        positions = list(range(len(labels)))
        if chart_type == "line":
            for item in series:
                axis.plot(labels, item["values"], marker="o", label=item["name"])
            axis.tick_params(axis="x", labelrotation=30)
            axis.set_xlabel(config.get("x_label") or "")
            axis.set_ylabel(config.get("y_label") or "")
            axis.legend()
        else:
            width = 0.8 / len(series)
            for series_index, item in enumerate(series):
                offsets = [
                    position + (series_index - (len(series) - 1) / 2) * width
                    for position in positions
                ]
                axis.bar(offsets, item["values"], width=width, label=item["name"])
            axis.set_xticks(positions, labels, rotation=30, ha="right")
            axis.set_xlabel(config.get("x_label") or "")
            axis.set_ylabel(config.get("y_label") or "")
            axis.legend()
        figure.tight_layout()
    axis.set_title(str(config.get("title") or "Graphique"))
    output = io.BytesIO()
    figure.savefig(output, format="png", bbox_inches="tight")
    plt.close(figure)
    return output.getvalue()


def _paragraph_text(value: Any) -> str:
    return escape(str(value)).replace("\n", "<br/>")


def build_pdf_export(
    *,
    question: str,
    answer: str,
    created_at: datetime,
    source_files: list[str],
    result_table: Any,
    chart_config: Any,
    conversation_turns: list[dict[str, Any]] | None = None,
    logo_path: Path | None = None,
) -> bytes:
    output = io.BytesIO()
    filename_text = ", ".join(source_files) or "Aucun fichier source"

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("ExportSans", 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        footer_text = f"{created_at:%d/%m/%Y %H:%M} · {filename_text[:150]}"
        canvas.drawString(16 * mm, 10 * mm, footer_text)
        canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, str(document.page))
        canvas.restoreState()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Rapport d’analyse",
    )
    styles = getSampleStyleSheet()
    for style_name in ("BodyText", "Heading1", "Heading2", "Heading3"):
        styles[style_name].fontName = "ExportSans" if style_name == "BodyText" else "ExportSans-Bold"
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        alignment=0,
        textColor=colors.HexColor("#0b1f3a"),
        fontName="ExportSans-Bold",
        fontSize=22,
        leading=28,
    )
    brand_style = ParagraphStyle(
        "ReportBrand",
        parent=styles["BodyText"],
        fontName="ExportSans-Bold",
        textColor=colors.HexColor("#2563EB"),
        fontSize=11,
        leading=14,
    )
    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["BodyText"],
        fontName="ExportSans",
        textColor=colors.HexColor("#0B1F3A"),
        leading=16,
        spaceAfter=9,
    )
    question_style = ParagraphStyle(
        "ReportQuestion",
        parent=body_style,
        textColor=colors.HexColor("#2563EB"),
        fontName="ExportSans-Bold",
        backColor=colors.HexColor("#EFF6FF"),
        borderPadding=9,
        spaceBefore=3,
        spaceAfter=12,
    )
    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=body_style,
        fontName="ExportSans",
        fontSize=7,
        leading=9,
        spaceAfter=0,
    )
    logo = logo_path or _logo_path
    story = [
        *([Image(str(logo), width=34 * mm, height=34 * mm, kind="proportional"), Spacer(1, 3 * mm)] if logo.is_file() else []),
        Paragraph("AI Data Analyst", brand_style),
        Paragraph("Rapport d’analyse", title_style),
        Spacer(1, 3 * mm),
        Paragraph(f"Analysé le {_paragraph_text(created_at.strftime('%d/%m/%Y à %H:%M'))}", body_style),
        Paragraph(f"Fichier{'' if len(source_files) == 1 else 's'} : {_paragraph_text(filename_text)}", body_style),
    ]

    turns = conversation_turns or [{
        "question": question,
        "answer": answer,
        "created_at": created_at,
        "result_table": result_table,
        "chart_config": chart_config,
    }]
    for index, turn in enumerate(turns):
        turn_question = str(turn.get("question") or question)
        turn_answer = str(turn.get("answer") or "")
        turn_date = turn.get("created_at")
        date_text = turn_date.strftime("%d/%m/%Y à %H:%M") if isinstance(turn_date, datetime) else ""
        story.extend([
            Spacer(1, 4 * mm),
            Paragraph(f"Question {index + 1}" if len(turns) > 1 else "Question", styles["Heading2"]),
            Paragraph(_paragraph_text(turn_question), question_style),
        ])
        if date_text:
            story.append(Paragraph(f"Réponse du {_paragraph_text(date_text)}", styles["Heading3"]))
        story.extend([
            Paragraph("Réponse de l’assistant", styles["Heading3"]),
            Paragraph(_paragraph_text(turn_answer), body_style),
        ])

        image_data = _chart_png(turn.get("chart_config"))
        if image_data:
            story.extend([
                Paragraph("Graphique", styles["Heading3"]),
                Image(io.BytesIO(image_data), width=170 * mm, height=95 * mm, kind="proportional"),
                Spacer(1, 3 * mm),
            ])

        columns, rows = _clean_table(turn.get("result_table"))
        if columns:
            table_data = [
                [Paragraph(_paragraph_text(value), table_cell_style) for value in columns],
                *[
                    [Paragraph(_paragraph_text(value if value is not None else ""), table_cell_style) for value in row]
                    for row in rows
                ],
            ]
            widths = [document.width / len(columns)] * len(columns)
            result_table_pdf = LongTable(table_data, colWidths=widths, repeatRows=1)
            result_table_pdf.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1F3A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFDBFE")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([Paragraph("Tableau de résultat", styles["Heading3"]), result_table_pdf])

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_filename(analysis_id: int, file_format: str) -> str:
    return f"analyse-{analysis_id}-resultat.{file_format}"


def pdf_export_filename(question: str, analysis_id: int, unique: bool = False) -> str:
    import re
    import unicodedata

    normalized = unicodedata.normalize("NFKD", question).encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"[^A-Za-z0-9]+", " ", normalized).strip().lower()
    normalized = re.sub(
        r"^(?:quel est le|quelle est la|quels sont les|quelles sont les|quel est|quelle est|quels sont|quelles sont|"
        r"comment|combien|pourquoi|donne moi|montre moi|peux tu|est ce que)\s+",
        "",
        normalized,
    )
    slug = re.sub(r"[^A-Za-z0-9]+", "-", normalized).strip("-")
    slug = slug[:48].strip("-") or "rapport"
    suffix = f"-{analysis_id}" if unique else ""
    return f"analyse-{slug}{suffix}.pdf"
