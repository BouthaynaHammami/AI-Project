import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import pandas as pd

ALLOWED_EXTENSIONS = {".csv", ".xls", ".xlsx", ".pdf", ".doc", ".docx", ".png", ".jpg", ".jpeg"}

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


def _read_legacy_word_document(file_path: Path) -> str:
    converter = shutil.which("antiword") or shutil.which("catdoc")
    if converter is not None:
        try:
            result = subprocess.run(
                [converter, str(file_path)],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=30,
            )
        except (subprocess.SubprocessError, OSError) as error:
            raise ValueError("The .doc file could not be converted to text.") from error
        return result.stdout.strip()

    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        raise ValueError(
            "Reading .doc files requires LibreOffice, antiword, or catdoc. "
            "You can also save the document as .docx."
        )

    try:
        with tempfile.TemporaryDirectory() as output_dir:
            subprocess.run(
                [office, "--headless", "--convert-to", "txt:Text", "--outdir", output_dir, str(file_path)],
                check=True,
                capture_output=True,
                timeout=30,
            )
            converted_path = Path(output_dir) / f"{file_path.stem}.txt"
            return converted_path.read_text(encoding="utf-8", errors="replace").strip()
    except (OSError, subprocess.SubprocessError) as error:
        raise ValueError("The .doc file could not be converted to text.") from error


def _ocr_image(image) -> str:
    try:
        import pytesseract

        return pytesseract.image_to_string(image, lang="fra+eng")
    except ImportError as error:
        raise ValueError("Image extraction requires the pytesseract Python package.") from error
    except Exception as error:
        if error.__class__.__name__ == "TesseractNotFoundError":
            raise ValueError(
                "Image text extraction requires Tesseract OCR and the French/English language data."
            ) from error
        raise ValueError("OCR could not read the image. Check that Tesseract has French and English data installed.") from error


def get_document_page_count(file_path: Path) -> int | None:
    extension = file_path.suffix.lower()
    if extension == ".pdf":
        import pymupdf

        with pymupdf.open(file_path) as document:
            return document.page_count

    if extension == ".docx":
        try:
            with ZipFile(file_path) as archive:
                properties = ElementTree.fromstring(archive.read("docProps/app.xml"))
            pages = properties.findtext(
                "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Pages"
            )
            return int(pages) if pages and int(pages) > 0 else None
        except (BadZipFile, KeyError, ValueError, ElementTree.ParseError):
            return None

    if extension == ".doc":
        office = shutil.which("soffice") or shutil.which("libreoffice")
        if office is None:
            return None
        try:
            with tempfile.TemporaryDirectory() as output_dir:
                subprocess.run(
                    [office, "--headless", "--convert-to", "pdf", "--outdir", output_dir, str(file_path)],
                    check=True,
                    capture_output=True,
                    timeout=30,
                )
                converted_pdf = Path(output_dir) / f"{file_path.stem}.pdf"
                return get_document_page_count(converted_pdf) if converted_pdf.exists() else None
        except (OSError, subprocess.SubprocessError):
            return None

    return None


def read_file_as_dataframe(file_path: Path) -> pd.DataFrame:
    """Read tabular files or extract document/image text into a DataFrame."""
    extension = file_path.suffix.lower()

    try:
        if extension == ".csv":
            try:
                dataframe = pd.read_csv(file_path, sep=None, engine="python", encoding="utf-8-sig")
            except UnicodeDecodeError:
                dataframe = pd.read_csv(file_path, sep=None, engine="python", encoding="cp1252")
        elif extension in {".xls", ".xlsx"}:
            dataframe = pd.read_excel(file_path)
        elif extension == ".pdf":
            import pymupdf
            from PIL import Image

            pages = []
            with pymupdf.open(file_path) as document:
                for page_number, page in enumerate(document, start=1):
                    text = page.get_text().strip()
                    if not text:
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                        image = Image.open(BytesIO(pixmap.tobytes("png")))
                        text = _ocr_image(image).strip()
                    pages.append({"page": page_number, "content": text})
            dataframe = pd.DataFrame(pages)
        elif extension == ".docx":
            from docx import Document

            document = Document(str(file_path))
            content = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                content.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                )
            dataframe = pd.DataFrame({"content": ["\n".join(content)]})
        elif extension == ".doc":
            dataframe = pd.DataFrame({"content": [_read_legacy_word_document(file_path)]})
        elif extension in {".png", ".jpg", ".jpeg"}:
            from PIL import Image

            with Image.open(file_path) as image:
                text = _ocr_image(image)
            dataframe = pd.DataFrame({"content": [text.strip()]})
        else:
            raise ValueError("Unsupported file format.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The uploaded file could not be read.") from error

    has_text_content = "content" not in dataframe.columns or dataframe["content"].fillna("").astype(str).str.strip().ne("").any()
    if (
        dataframe.empty
        or not has_text_content
        or dataframe.replace(r"^\s*$", pd.NA, regex=True).dropna(how="all").empty
    ):
        raise ValueError("The uploaded file contains no readable data or text.")

    return dataframe


def save_and_read_file(filename: str, content: bytes):
    """
    Save an uploaded file and read it with Pandas.

    Args:
        filename: Original filename.
        content: Uploaded file content.

    Returns:
        Tuple containing the saved file path and Pandas DataFrame.

    Raises:
        ValueError: If the file format is unsupported or invalid.
    """

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file format. Use CSV, XLS/XLSX, PDF, DOC/DOCX, PNG, or JPG."
        )

    file_path = UPLOAD_DIR / f"{uuid4().hex}{extension}"
    file_path.write_bytes(content)

    try:
        dataframe = read_file_as_dataframe(file_path)
    except ValueError:
        file_path.unlink(missing_ok=True)
        raise

    return file_path, dataframe